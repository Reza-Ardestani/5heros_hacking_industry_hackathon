#!/usr/bin/env python3
"""Backfill the Calgary disruption database and export JSON + Excel from it.

1. Downloads every City dataset for the window (default six months) into the
   SQLite store (data/db/disruptions.sqlite; BB_DB_PATH overrides). Each download
   is logged in fetch_runs with its URL, SHA256 and row counts.
2. Exports the database contents to data/analysis/ (JSON files + workbook). The
   JSON exports are committed so a fresh checkout can seed its database offline.

Parsing/classification lives in backend/app/domain/disruption_rules.py and
backend/app/application/disruption_ingest.py; this script only orchestrates.

Run:  uv run --no-project --with openpyxl --with tzdata python scripts/build_disruption_dataset.py
      add --export-only to re-export without downloading.
"""

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.application import disruptions
from app.application.disruption_collector import Collector
from app.application.disruption_ingest import CAMERA_RADIUS_M, SIGNAL_RADIUS_M
from app.bootstrap import configure_services
from app.infra import city_open_data
from app.infra.disruption_store import Store
from app.infra.raw_sink import JsonRawSink

configure_services()


def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--months", type=int, default=6)
    ap.add_argument("--out", type=Path, default=ROOT / "data" / "analysis")
    ap.add_argument("--raw", type=Path, default=ROOT / "data" / "raw" / "calgary_disruptions")
    ap.add_argument("--db", type=Path, default=None, help="SQLite path (default data/db/...)")
    ap.add_argument("--export-only", action="store_true")
    args = ap.parse_args()
    store = Store(args.db)
    if not args.export_only:
        print(f"Backfilling {args.months} months into {store.path} …")
        report = Collector(
            store, fetcher=city_open_data.fetch, raw_sink=JsonRawSink(args.raw)
        ).backfill(args.months)
        for key, r in report.items():
            print(f"  {key}: {r}")
        failed = [k for k, r in report.items() if r["status"] != "ok"]
        if failed:
            print(f"WARNING: failed datasets {failed}; previous rows are kept", file=sys.stderr)
    export(store, args.out)
    print(json.dumps(store.stats(), indent=1))


def export(store, out):
    disruptions.use_store(store)
    data = disruptions._load()
    full = disruptions.full_summary()
    manifest = full["manifest"]
    incidents = [_export_incident(r) for r in data["incidents"]]
    closures = data["closures"]
    ref = data["reference"]
    projects = [
        {
            "title": r.get("title"),
            "type": r.get("projecttype"),
            "budget": r.get("budget"),
            "completion": r.get("completion"),
            "url": r.get("url"),
            "description": " ".join((r.get("description") or "").split()),
            "longitude": ((r.get("point") or {}).get("coordinates") or [None, None])[0],
            "latitude": ((r.get("point") or {}).get("coordinates") or [None, None])[1],
        }
        for r in store.get_reference("projects")[0] or []
    ]
    latest_tt = {}
    for t in store.travel_times(limit=100000):
        latest_tt.setdefault(t["segment"], t)
    travel_times = [
        {"corridor": t["corridor"], "segment": t["segment"],
         "travel_time_min": t["travel_time_min"], "last_update_local": t["city_updated_local"]}
        for t in latest_tt.values()
    ]  # fmt: skip
    current = store.get_reference("current_incidents")[0] or []
    summary = {k: v for k, v in full.items() if k != "manifest"}
    summary["data_quality"]["database"] = store.stats()
    out.mkdir(parents=True, exist_ok=True)
    write_json(out / "calgary_incidents_6mo.json", {"manifest": manifest, "records": incidents})
    write_json(out / "calgary_closures.json", {"manifest": manifest, "records": closures})
    write_json(
        out / "calgary_context_layers.json",
        {
            "manifest": manifest,
            "current_incidents": current,
            "construction_projects": projects,
            "travel_times_snapshot": travel_times,
            "cameras": ref.cameras,
            "signals": ref.signals,
            "volumes_2024": [{k: v for k, v in x.items() if k != "coords"} for x in ref.volumes],
        },
    )
    write_json(out / "calgary_disruption_summary.json", {"manifest": manifest, **summary})
    write_json(out / "sources_manifest.json", manifest)
    write_excel(
        out / "calgary_disruptions_analysis.xlsx",
        manifest, incidents, closures, current, projects, travel_times, summary,
    )  # fmt: skip
    print(f"Exported {len(incidents)} incidents, {len(closures)} closures to {out}")


def _export_incident(r):
    # Database bookkeeping fields stay in the DB; exports keep the parsed record + uid.
    drop = {"stored_source", "first_seen_utc"}
    return {k: v for k, v in r.items() if k not in drop}


def write_json(path, obj):
    path.write_text(
        json.dumps(obj, indent=1, ensure_ascii=False, default=str) + "\n", encoding="utf-8"
    )


def write_excel(path, manifest, incidents, closures, current, projects, travel_times, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="1F4E5F"), Font(bold=True, color="FFFFFF")

    def sheet(title, rows, cols=None, note=None):
        ws = wb.create_sheet(title[:31])
        cols = cols or (list(rows[0].keys()) if rows else ["(no rows)"])
        r0 = 1
        if note:
            ws.cell(1, 1, note).font = Font(italic=True, color="555555")
            r0 = 2
        for j, c in enumerate(cols, 1):
            cell = ws.cell(r0, j, c)
            cell.fill, cell.font = head_fill, head_font
        for i, row in enumerate(rows, r0 + 1):
            for j, c in enumerate(cols, 1):
                v = row.get(c)
                ws.cell(i, j, ", ".join(map(str, v)) if isinstance(v, list) else v)
        ws.freeze_panes = ws.cell(r0 + 1, 1)
        if rows:
            ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0 + len(rows)}"
        for j, c in enumerate(cols, 1):
            width = max([len(str(c))] + [len(str(r.get(c) or "")) for r in rows[:300]])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 60)
        return ws

    ws = wb.active
    ws.title = "README"
    h = summary["headline"]
    lines = [
        ("Calgary traffic disruptions — analysis workbook", True),
        (
            f"Exported from {manifest['database']}. Latest download {manifest['retrieved_at_utc']} UTC. Window: {manifest['window']['start_local'][:10]} to export ({h['window_days']} days).",
            False,
        ),
        (
            f"Discovered via {manifest['discovered_from']} → City of Calgary Open Data feeds (see Sources).",
            False,
        ),
        (manifest["attribution"] + ". " + manifest["license_url"], False),
        ("", False),
        ("Headline", True),
        (
            f"Incidents (deduplicated): {h['incidents']}  |  per day: {h['incidents_per_day']}  |  lane-blocking: {h['lane_blocking_pct']}%",
            False,
        ),
        (
            f"At a signalized intersection (≤{SIGNAL_RADIUS_M} m): {h['at_signalized_intersection_pct']}%  |  within {CAMERA_RADIUS_M} m of a traffic camera: {h['within_500m_of_camera_pct']}%",
            False,
        ),
        (
            f"Closures in detour feed: {h['closures_in_feed']} ({h['closures_overlapping_window']} overlap the window)",
            False,
        ),
        ("", False),
        ("How to use for Bottleneck Busters", True),
        (
            "• Corridors / Intersections / Signal hotspots rank where reported disruptions recur — candidates for a corridor study, NOT measured congestion.",
            False,
        ),
        (
            "• incidents_per_10k_daily_veh divides by City 2024 segment volume (median of same-road segments ≤150 m) — a crude exposure normalisation; volume is a year-2024 estimate.",
            False,
        ),
        (
            "• Signal-fault and lane-blocking columns map to the planner's incident features (lane impact, direction, location/time).",
            False,
        ),
        (
            "• Closures overlapping a corridor indicate construction-driven capacity loss in the same period.",
            False,
        ),
        ("", False),
        ("Limitations", True),
        (
            "• Incident archive is unofficial, camera-imagery based; gaps possible; not traffic flow, delay or exposure-adjusted crash risk.",
            False,
        ),
        (
            "• Categories are rule-based text classification of the City's free-text description (see Category rules). 'Traffic incident (unspecified)' carries no cause.",
            False,
        ),
        (
            "• update_lag_min = last update minus start; it is NOT clearance time or incident duration.",
            False,
        ),
        (
            "• Closure feed lists current and scheduled closures only; closures that ended before retrieval are not retained, so historical closure coverage is incomplete.",
            False,
        ),
        ("• Travel times and current incidents are point-in-time snapshots at retrieval.", False),
        (
            "• Road/corridor names are parsed from free text; spelling variants may split a corridor.",
            False,
        ),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=False)
    ws.column_dimensions["A"].width = 150

    src_rows = [{"key": k, **dict(v.items())} for k, v in manifest["sources"].items()]
    sheet("Sources", src_rows)
    hl = [{"metric": k, "value": json.dumps(v) if isinstance(v, dict) else v} for k, v in h.items()]
    hl += [
        {"metric": k, "value": json.dumps(v) if isinstance(v, (list, dict)) else v}
        for k, v in summary["data_quality"].items()
    ]
    sheet("Headline & quality", hl)
    sheet("By category", summary["by_category"])
    sheet("Monthly trend", summary["monthly"])
    sheet(
        "Hour x weekday",
        summary["hour_by_weekday"],
        note="Incident start counts, local time (America/Edmonton)",
    )
    sheet("By period", summary["by_period"])
    sheet("By quadrant", summary["by_quadrant"])
    sheet("Lane impact", summary["by_lane_impact"])
    sheet("Top corridors", summary["top_corridors"])
    sheet("Top intersections", summary["top_intersections"])
    sheet(
        "Signal hotspots",
        summary["signal_hotspots"],
        note=f"Incidents within {SIGNAL_RADIUS_M} m of a City traffic signal",
    )
    sheet(
        "Camera hotspots",
        summary["camera_hotspots"],
        note=f"Incidents within {CAMERA_RADIUS_M} m of a traffic camera",
    )
    sheet("Incidents", incidents)
    sheet("Closures", closures)
    sheet(
        "Closure summary",
        summary["closures_by_type"]
        + [{"closure_type": "—"}]
        + [
            {"closure_type": f"status: {x['status']}", "closures": x["closures"]}
            for x in summary["closures_by_status"]
        ],
    )
    sheet("Current incidents", current, note="Snapshot at retrieval")
    sheet("Construction projects", projects)
    sheet("Travel times", travel_times, note="Snapshot at retrieval")
    sheet(
        "Category rules",
        summary["category_rules"],
        note="First matching regex (case-insensitive) on description wins",
    )
    wb.save(path)


if __name__ == "__main__":
    main()
